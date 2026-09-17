import io
import os
import shutil
import tempfile
import logging

logger = logging.getLogger(__name__)


def _find_libreoffice_binary() -> str | None:
    """Finds libreoffice or soffice executable path if available."""
    for cmd in ["libreoffice", "soffice", "libreoffice7.6", "libreoffice7.5"]:
        path = shutil.which(cmd)
        if path:
            return path
    return None


def _convert_via_docx2pdf(docx_bytes: bytes) -> bytes:
    """Windows path: docx2pdf via MS Word COM."""
    from docx2pdf import convert
    import pythoncom

    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, "input.docx")
        pdf_path = os.path.join(tmpdir, "input.pdf")

        with open(docx_path, "wb") as f:
            f.write(docx_bytes)

        # Initialize COM for this thread (required when called from thread pool)
        pythoncom.CoInitialize()
        try:
            convert(docx_path, pdf_path)
        finally:
            pythoncom.CoUninitialize()

        with open(pdf_path, "rb") as f:
            return f.read()


def _convert_via_libreoffice(docx_bytes: bytes) -> bytes:
    """Cross-platform fallback: LibreOffice headless CLI."""
    import subprocess

    cmd = _find_libreoffice_binary()
    if not cmd:
        raise FileNotFoundError("LibreOffice binary not found in PATH")

    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, "input.docx")
        with open(docx_path, "wb") as f:
            f.write(docx_bytes)

        result = subprocess.run(
            [
                cmd, "--headless", "--convert-to", "pdf",
                "--outdir", tmpdir, docx_path,
            ],
            capture_output=True,
            timeout=60,
        )
        if result.returncode != 0:
            raise RuntimeError(f"LibreOffice conversion failed: {result.stderr.decode()}")

        pdf_path = os.path.join(tmpdir, "input.pdf")
        if not os.path.exists(pdf_path):
            raise RuntimeError("LibreOffice conversion produced no output PDF file")

        with open(pdf_path, "rb") as f:
            return f.read()


def _convert_via_reportlab(docx_bytes: bytes) -> bytes:
    """
    Pure-Python fallback when docx2pdf and LibreOffice CLI binaries are unavailable.
    Parses docx paragraphs/tables using python-docx and builds a clean PDF via ReportLab.
    """
    import docx
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    doc = docx.Document(io.BytesIO(docx_bytes))
    buffer = io.BytesIO()
    pdf_doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    normal_style = styles['Normal']

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=6,
    )
    heading_style = ParagraphStyle(
        'DocHeading',
        parent=styles['Heading2'],
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#1e40af"),
        spaceBefore=8,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=normal_style,
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#334155"),
        spaceAfter=3,
    )

    story = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style_name = p.style.name.lower() if p.style else ""
        if "heading 1" in style_name or "title" in style_name:
            story.append(Paragraph(text, title_style))
        elif "heading" in style_name:
            story.append(Paragraph(text, heading_style))
        else:
            story.append(Paragraph(text, body_style))

    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [Paragraph(cell.text.strip(), body_style) for cell in row.cells]
            table_data.append(row_data)
        if table_data:
            t = Table(table_data)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#f1f5f9')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#0f172a')),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                ('TOPPADDING', (0,0), (-1,-1), 4),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
            ]))
            story.append(t)
            story.append(Spacer(1, 6))

    if not story:
        story.append(Paragraph("Resume Document", title_style))

    pdf_doc.build(story)
    return buffer.getvalue()


def convert_docx_to_pdf(docx_bytes: bytes) -> bytes:
    """
    Public interface.
    1. Try docx2pdf (Windows/MS Word) if on Windows.
    2. Try LibreOffice/soffice if binary exists on system PATH.
    3. Fall back to pure-Python ReportLab conversion (works on any OS/Render without binaries).
    """
    import sys

    if sys.platform == "win32":
        try:
            import docx2pdf  # noqa: F401
            return _convert_via_docx2pdf(docx_bytes)
        except ImportError:
            logger.warning("docx2pdf not available on Windows, trying alternatives...")
        except Exception as e:
            logger.warning("docx2pdf failed (%s), trying alternatives...", e)

    # Check for LibreOffice binary
    if _find_libreoffice_binary():
        try:
            logger.info("Converting DOCX to PDF using LibreOffice")
            return _convert_via_libreoffice(docx_bytes)
        except Exception as e:
            logger.warning("LibreOffice conversion failed (%s), trying pure-Python fallback...", e)

    # Fall back to ReportLab pure Python conversion
    logger.info("Using pure-Python ReportLab PDF conversion fallback")
    try:
        return _convert_via_reportlab(docx_bytes)
    except Exception as e:
        logger.error("ReportLab conversion failed: %s", e)
        raise RuntimeError(f"PDF conversion failed: {e}")

